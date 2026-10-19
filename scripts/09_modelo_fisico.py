"""09 - Modelo fisico-operacional (camada 1 da arquitetura P1): necessidade de reator -> manobras do fixo, taps e manobras residuais do VSR.

Fisica. A sobra de reativo de uma linha e Qc*(1 - (P/SIL)^2): o carregamento capacitivo menos o que a propria linha consome (I^2 X) ao transportar P.
A compensacao permanentemente ligada cobre a necessidade com a linha carregada (percentil 95 do carregamento u = |P|); a compensacao
manobravel cobre a variacao W entre o p95 e o p5 do carregamento. Variavel normalizada, sem corte:
    r(t) = (u_hi^2 - u(t)^2) / (u_hi^2 - u_lo^2)        (r = 0 no p95 do carregamento; r = 1 no p5)
u = |fluxo verificado N-NE| (Campo) ou |carga liquida do NE| (Centro Urbano, Polo eolico), que carrega as linhas que alimentam a area.

Arranjo fixo: N degraus manobraveis de W/N. Cada degrau entra quando r passa do meio do degrau + histerese/2 e sai abaixo do meio - histerese/2
(histerese = eta x tamanho do degrau: faixa de tensao que o operador tolera antes de manobrar).
VSR, arranjo 1:1: VSR de tamanho W substitui so o banco manobravel; faixa [m, 1] em r; sai pelo disjuntor quando r < m/2 - histerese/2.
VSR, arranjo 2:1 (CIGRE A2-211: 1 VSR no lugar de 2 fixos): VSR de tamanho W/(1-m) substitui o banco manobravel e uma fatia m*W/(1-m) da
compensacao permanente; segue r em [0, 1] por taps e so sai quando r < -m/(2(1-m)) (carregamento extremo). No arranjo fixo equivalente, a
fatia substituida segue a mesma regra, entao manobras evitadas = manobras do banco manobravel.
Taps: 33 posicoes na faixa do VSR (Siemens/Amprion); o comutador move quando o erro passa de 0,75 passo. Resolucao de 30 min (DESSEM) e
1 h (fluxo N-NE): taps/dia sao piso. Volatilidade (P4): x_k = perfil + k (x - perfil); dimensionamento (u_lo, u_hi) fixo em k = 1.

Saidas: processed/corredor_NNE_linhas.csv, modelo_fisico_taxas.csv, modelo_fisico_mensal.csv, modelo_fisico_validacao.csv,
modelo_fisico_camada1.json; abas Corredor_NNE e ModeloFisico_* na planilha; linhas P79+ de Parametros_Modelo; LEIA-ME.
Uso: python 09_modelo_fisico.py <pasta_base_PoliVSR-dados> <planilha.xlsx>   (rodar depois do 05 e do 08)"""
import sys, glob, json, time, numpy as np, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
BASE, XL = sys.argv[1], sys.argv[2]; R = BASE + '/raw/ons/'; P = BASE + '/processed/'
ETA0, M0, K0, N0, ARR0, T_POS, DB_TAP, MVAR_UNID = 0.3, 0.5, 1.0, 1, '2:1', 33, 0.75, 150.0   # 150 Mvar = mediana do reator 500 kV (cadastro ONS)
NOME = {'urbano': 'Centro Urbano', 'campo': 'Campo (interligação N–NE)', 'polo_eolico': 'Polo eólico'}
t0 = time.time()

# ======================= 1. Corredor N-NE =======================
L = pd.read_csv(R + 'ONS__linha_transmissao__LINHA_TRANSMISSAO.csv', sep=';')
for c in ['id_subsistema_terminalde', 'id_subsistema_terminalpara', 'nom_tipolinha', 'nom_linhadetransmissao']: L[c] = L[c].astype(str).str.strip()
act = L[L.dat_desativacao.isna() & L.dat_entradaoperacao.notna()]
C = act[act.id_subsistema_terminalde.isin(['N', 'NE']) & act.id_subsistema_terminalpara.isin(['N', 'NE'])
        & (act.id_subsistema_terminalde != act.id_subsistema_terminalpara)].copy()
C['SIL_MW'] = 100 * np.sqrt((C.val_shunt / 100) / (C.val_reatancia / 100))   # X e B em % na base 100 MVA (convencao ANAREDE)
C['Mvar_km'] = C.val_shunt / C.val_comprimento
COR = C[['nom_linhadetransmissao', 'val_niveltensao_kv', 'val_comprimento', 'val_reatancia', 'val_shunt', 'SIL_MW', 'Mvar_km',
         'val_capacoperlongasemlimit', 'dat_entradaoperacao']].rename(columns={
    'nom_linhadetransmissao': 'linha', 'val_niveltensao_kv': 'kV', 'val_comprimento': 'km', 'val_reatancia': 'X_pct_100MVA',
    'val_shunt': 'Qc_Mvar', 'val_capacoperlongasemlimit': 'capacidade_longa_A', 'dat_entradaoperacao': 'entrada_operacao'}).sort_values(['kV', 'linha'])
COR['linha'] = COR.linha.str.replace(r'\s+', ' ', regex=True)
SIL, QC = COR.SIL_MW.sum(), COR.Qc_Mvar.sum()
tot = {'linha': 'TOTAL', 'kV': np.nan, 'km': COR.km.sum(), 'X_pct_100MVA': np.nan, 'Qc_Mvar': QC, 'SIL_MW': SIL, 'Mvar_km': QC / COR.km.sum(),
       'capacidade_longa_A': np.nan, 'entrada_operacao': f'{len(COR)} circuitos'}
COR = pd.concat([COR, pd.DataFrame([tot])], ignore_index=True)
for c in ('km', 'Qc_Mvar', 'SIL_MW', 'Mvar_km'): COR[c] = COR[c].astype(float).round(3)
COR.to_csv(P + 'corredor_NNE_linhas.csv', index=False, sep=';', encoding='utf-8-sig')
print(f'Corredor N–NE: {len(COR) - 1} circuitos, {COR.km.iloc[-1]:.0f} km, Qc {QC:.0f} Mvar, SIL {SIL:.0f} MW')

# ======================= 2. Series =======================
d = pd.read_csv(glob.glob(R + '*balanco_dessem_detalhe__CONSOLIDADO*.csv')[0], sep=';')
d['date'] = pd.to_datetime(d.din_programacaodia); d['slot'] = d.num_patamar.astype(int)
liq = lambda g: g.val_demanda - g.val_ger_mmgd - g.val_ger_fotovoltaica - g.val_ger_eolica
gdl = lambda g: g.val_demanda - g.val_ger_mmgd - g.val_ger_fotovoltaica
def dser(sub, f):
    g = d[d.cod_subsistema == sub].drop_duplicates(['date', 'slot']).sort_values(['date', 'slot'])
    c = g.groupby('date').slot.nunique(); g = g[g.date.isin(c[c == 48].index)]
    return pd.DataFrame({'date': g.date.values, 'slot': g.slot.values, 'x': f(g).values})
fl = pd.read_csv(P + 'fluxo_NNE_horario.csv', sep=';', encoding='utf-8-sig')
fl['t'] = pd.to_datetime(fl.t); fl['date'] = fl.t.dt.normalize(); fl['slot'] = fl.t.dt.hour + 1
c = fl.groupby('date').slot.nunique(); fl = fl[fl.date.isin(c[c == 24].index)].sort_values('t')
SER = {'urbano': (dser('NE', gdl), 0.5, 'NE: demanda − MMGD − FV (ONS DESSEM, programado)'),
       'campo': (pd.DataFrame({'date': fl.date.values, 'slot': fl.slot.values, 'x': fl.fluxo_N_para_NE_MW.values}), 1.0,
                 'Fluxo verificado N–NE (ONS, intercâmbio entre subsistemas)'),
       'polo_eolico': (dser('NE', liq), 0.5, 'NE: carga líquida (ONS DESSEM, programado)')}
def perfil(df):
    mes = df.date.dt.to_period('M'); td = np.where(df.date.dt.dayofweek < 5, 'u', 'f')
    return df.groupby([mes, td, df.slot]).x.transform('mean').values

# ======================= 3. Simulacao =======================
def sim(r, N, arr, m, eta, T=T_POS, db=DB_TAP):
    r = [float(v) for v in r]; n = len(r)
    th = [(i + 0.5) / N for i in range(N)]; hs = eta / N
    if arr == '2:1': a, b, Bm = 0.0, 1.0, m / (1 - m)
    else: a, b, Bm = m, 1.0, m
    thr = a - Bm / 2; hv = eta * Bm; st = (b - a) / (T - 1)
    on = [False] * N; base_on = True; con = True; p = 0
    o_sw = np.zeros(n); o_base = np.zeros(n); res = np.zeros(n); taps = np.zeros(n); lf_ = np.zeros(n); lv_ = np.zeros(n)
    for t, v in enumerate(r):
        k_on = 0
        for i in range(N):
            if not on[i] and v > th[i] + hs / 2: on[i] = True; o_sw[t] += 1
            elif on[i] and v < th[i] - hs / 2: on[i] = False; o_sw[t] += 1
            k_on += on[i]
        lf = k_on / N
        if arr == '2:1':
            if not base_on and v > thr + hv / 2: base_on = True; o_base[t] += 1
            elif base_on and v < thr - hv / 2: base_on = False; o_base[t] += 1
            if not base_on: lf -= Bm
        if not con and v > thr + hv / 2: con = True; p = 0; res[t] += 1
        elif con and v < thr - hv / 2: con = False; res[t] += 1
        if con:
            tgt = min(max(v, a), b)
            if abs(tgt - (a + p * st)) > db * st:
                q = int(round((tgt - a) / st)); taps[t] += abs(q - p); p = q
            lv = a + p * st
        else:
            lv = a - Bm
        lf_[t] = lf; lv_[t] = lv
    return dict(o_sw=o_sw, o_base=o_base, res=res, taps=taps, lf=lf_, lv=lv_)

def resumo(out, r, nd, N, arr):
    fixo = out['o_sw'] + out['o_base']
    return dict(dias=nd, fixo_ops_dia_total=fixo.sum() / nd, fixo_ops_dia_por_disjuntor_manobravel=out['o_sw'].sum() / nd / N,
                fatia_fixa_substituida_ops_dia=out['o_base'].sum() / nd, vsr_manobras_residuais_dia=out['res'].sum() / nd,
                manobras_evitadas_dia=(fixo.sum() - out['res'].sum()) / nd, vsr_taps_dia=out['taps'].sum() / nd,
                desajuste_fixo=np.abs(r - out['lf']).mean(), desajuste_vsr=np.abs(r - out['lv']).mean())

rows, mens, tip, val_extra, prep = [], [], {}, {}, {}
g0 = lambda rows_, cen_, N_: next(x['fixo_ops_dia_por_disjuntor_manobravel'] for x in rows_ if x['id'] == f'{cen_}|2:1|N{N_}|m{M0}|k{K0}|eta{ETA0}')
for cen, (df, step, desc) in SER.items():
    u0 = np.abs(df.x.values); prep[cen] = dict(lo=np.percentile(u0, 5), hi=np.percentile(u0, 95))
W_campo = QC * (prep['campo']['hi'] ** 2 - prep['campo']['lo'] ** 2) / SIL ** 2          # variacao da sobra de reativo do corredor, Mvar
N_campo = max(1, int(round(W_campo / MVAR_UNID)))
N_BASE = {'urbano': 1, 'campo': N_campo, 'polo_eolico': 1}
print(f'Campo: variação da sobra de reativo W = {W_campo:.0f} Mvar -> {N_campo} degraus de ~{MVAR_UNID:.0f} Mvar')
for cen, (df, step, desc) in SER.items():
    df = df.sort_values(['date', 'slot']).reset_index(drop=True)
    prof = perfil(df); lo, hi = prep[cen]['lo'], prep[cen]['hi']
    nd = df.date.nunique(); day = df.date.values
    prep[cen].update(dias=nd, inicio=str(df.date.min().date()), fim=str(df.date.max().date()), resolucao_h=step, definicao=desc, N_base=N_BASE[cen])
    for k in (0.5, 0.75, 1.0, 1.25, 1.5, 2.0):
        u = np.abs(prof + k * (df.x.values - prof)); r = (hi ** 2 - u ** 2) / (hi ** 2 - lo ** 2)
        for arr in ('1:1', '2:1'):
            for N in sorted({1, 2, 4, 6, N_BASE[cen]}):
                for m in (0.2, 0.3, 0.4, 0.5, 0.6):
                    etas = (0.1, 0.2, 0.3, 0.5, 0.75) if (k == K0 and N == N_BASE[cen] and m == M0) else (ETA0,)
                    for eta in etas:
                        out = sim(r, N, arr, m, eta)
                        rows.append(dict(id=f'{cen}|{arr}|N{N}|m{m}|k{k}|eta{eta}', cenario=cen, arranjo=arr, N=N, m=m, k=k, eta=eta, **resumo(out, r, nd, N, arr)))
                        if k == K0 and N == N_BASE[cen] and m == M0 and eta == ETA0 and arr == ARR0:
                            base = out; rbase = r; NB = N
    # ---- mensal e dia tipico (caso-base) ----
    fixo = base['o_sw'] + base['o_base']
    dfm = pd.DataFrame({'date': day, 'fixo': fixo, 'sw': base['o_sw'] / NB, 'res': base['res'], 'taps': base['taps']})
    dd = dfm.groupby('date').sum(); dd['mes'] = pd.to_datetime(dd.index).to_period('M').astype(str)
    mm = dd.groupby('mes').agg(dias=('fixo', 'size'), fixo_ops_dia_total=('fixo', 'mean'), ops_dia_por_disjuntor_manobravel=('sw', 'mean'),
                               vsr_manobras_residuais_dia=('res', 'mean'), vsr_taps_dia=('taps', 'mean')).reset_index()
    mm.insert(0, 'cenario', cen); mm.insert(1, 'N', NB); mens.append(mm)
    dts = pd.to_datetime(dd.index)
    verao = dd[(dts.month.isin([12, 1, 2, 3])) & (dts.dayofweek < 5)]
    val_extra[cen] = dict(verao_util=verao.sw.mean(), anual=dd.sw.mean(), n_verao=len(verao),
                          umido=dd[dts.month.isin([2, 3, 4, 5])].sw.mean(), seco=dd[dts.month.isin([8, 9, 10])].sw.mean())
    if cen == 'campo':
        u_mid = np.sqrt((hi ** 2 + lo ** 2) / 2); ua = np.abs(df.x.values); mo = df.date.dt.month.values
        perto = lambda msk_: 100 * (np.abs(ua[msk_] - u_mid) < 0.15 * u_mid).mean()
        val_extra[cen].update(u_mid=u_mid, perto_umido=perto(np.isin(mo, [2, 3, 4, 5])), perto_seco=perto(np.isin(mo, [8, 9, 10])))
    # dia tipico: dia util com 2 manobras do banco manobravel cuja amplitude de r esta mais perto da mediana
    rr = pd.DataFrame({'date': day, 'r': rbase}).groupby('date').r.agg(lambda v: v.max() - v.min())
    cand = dd[(dd.sw * NB >= 2) & (dts.dayofweek < 5)].index
    if len(cand) == 0: cand = dd.index
    dsel = (rr.loc[cand] - rr.median()).abs().idxmin()
    msk = day == np.datetime64(dsel)
    tip[cen] = dict(data=str(pd.Timestamp(dsel).date()), hora=[(s - 1) * step for s in df.slot.values[msk]],
                    r=np.round(rbase[msk], 4).tolist(), nivel_fixo=np.round(base['lf'][msk], 4).tolist(), nivel_vsr=np.round(base['lv'][msk], 4).tolist(),
                    manobras_fixo=int(fixo[msk].sum()), manobras_vsr=int(base['res'][msk].sum()), taps_vsr=int(base['taps'][msk].sum()))
    print(f'{cen}: {nd} dias, u p5 {lo:.0f} p95 {hi:.0f} MW | base 2:1 N={NB} m=0,5: {base["o_sw"].sum() / nd / NB:.2f} op/dia por disjuntor manobrável, VSR resid {base["res"].sum() / nd:.2f}, taps {base["taps"].sum() / nd:.1f}/dia ({time.time() - t0:.0f}s)')
    for N_ in sorted({1, 2, 4, 6, N_BASE[cen]}):
        print(f'    N={N_}: por disjuntor {g0(rows, cen, N_):.2f} op/dia')
TX = pd.DataFrame(rows)
for c in TX.columns[8:]: TX[c] = TX[c].astype(float).round(4)
MM = pd.concat(mens, ignore_index=True)
for c in MM.columns[4:]: MM[c] = MM[c].astype(float).round(3)
TX.to_csv(P + 'modelo_fisico_taxas.csv', index=False, sep=';', encoding='utf-8-sig')
MM.to_csv(P + 'modelo_fisico_mensal.csv', index=False, sep=';', encoding='utf-8-sig')

# ======================= 4. Validacao =======================
g = lambda cen, arr, N, m, k, eta, col: float(TX.loc[TX.id == f'{cen}|{arr}|N{N}|m{m}|k{k}|eta{eta}', col].iloc[0])
fm = pd.read_csv(P + 'fluxo_NNE_mensal.csv', sep=';', encoding='utf-8-sig')
inv_um = fm[fm.mes.str[5:7].isin(['02', '03', '04', '05'])].inversoes_por_dia.mean()
inv_se = fm[fm.mes.str[5:7].isin(['08', '09', '10'])].inversoes_por_dia.mean()
cl, ch = prep['campo']['lo'] / SIL, prep['campo']['hi'] / SIL
n_lo, Wc = 1 - ch ** 2, ch ** 2 - cl ** 2
m_max = n_lo / (n_lo + Wc)
V = []
for cen in SER:
    V.append(dict(id=f'{cen}_verao', cenario=cen, verificacao=f'Manobras/dia por disjuntor manobrável (N = {N_BASE[cen]}) – dia útil de verão (dez–mar)', modelo=val_extra[cen]['verao_util'],
                  referencia='≈ 2 (1 ciclo/dia por reator chaveado)', fonte='ONS – PAR/PEL 2025, p. 63–65 (~140 reatores de barra chaveados por dia útil de verão)',
                  leitura='Dentro da faixa observada (0,5–2 op/dia)' if 0.4 <= val_extra[cen]['verao_util'] <= 2.2 else 'Fora da faixa observada'))
    V.append(dict(id=f'{cen}_anual', cenario=cen, verificacao=f'Manobras/dia por disjuntor manobrável (N = {N_BASE[cen]}) – média do período', modelo=val_extra[cen]['anual'],
                  referencia='≈ 2 num reator em manobra diária (PAR/PEL); 0,48 num reator de barra 400 kV (175 op/ano)', fonte='ONS – PAR/PEL 2025; Jeyaraj & Habtay 2011, via Oliveira (UFCG 2015)',
                  leitura='Entre a média da frota e o reator em manobra diária' if 0.4 <= val_extra[cen]['anual'] <= 2.2 else 'Fora da faixa observada'))
V.append(dict(id='frota_media', cenario='todos', verificacao='Média da frota manobrável (cálculo sobre as referências, não saída do modelo)', modelo=140 * 2 / 530,
              referencia='0,48 op/dia num reator de barra 400 kV (Jeyaraj & Habtay 2011)', fonte='ONS – PAR/PEL 2025 (~140 reatores chaveados por dia útil de verão, 2 operações cada) ÷ CIGRÉ A3-10747 (~530 manobráveis)',
              leitura='Os cenários representam um reator em manobra diária (1,6–2 op/dia); a média da frota (~0,5 op/dia) é a base para os números nacionais do dashboard'))
ve = val_extra['campo']
V.append(dict(id='campo_sazonal', cenario='campo', verificacao='Campo – manobras/dia por disjuntor no período úmido (fev–mai); leitura física, não teste',
              modelo=ve['umido'], referencia=f'Seco (ago–out): {ve["seco"]:.2f}/dia. O banco manobrável troca de estado quando |P| cruza ~{ve["u_mid"]:,.0f} MW ({100 * ve["u_mid"] / SIL:.0f}% do SIL); '
                                             f'|P| fica a ±15% desse nível em {ve["perto_umido"]:.0f}% das horas no úmido e {ve["perto_seco"]:.0f}% no seco. Inversões de sentido (|P| ≈ 0) mantêm os reatores ligados.',
              fonte='ONS – intercâmbio entre subsistemas (abas Fluxo_NNE_*) + cadastro de linhas', leitura='As manobras do Campo seguem as oscilações do fluxo em torno do nível médio, não as inversões'))
V.append(dict(id='campo_N', cenario='campo', verificacao='Campo – degraus manobráveis do corredor (variação da sobra de reativo ÷ 150 Mvar)', modelo=N_campo,
              referencia=f'W = Qc·(u_p95² − u_p5²)/SIL² = {W_campo:,.0f} Mvar; reator 500 kV mediano = 150 Mvar', fonte='Cadastro ONS de linhas e de reatores + fluxo verificado N–NE',
              leitura='Caso-base do Campo usa esse N; Centro Urbano e Polo eólico usam N = 1 (sem SIL de referência)'))
V.append(dict(id='campo_m_max', cenario='campo', verificacao='Mínimo de regulação máximo viável no arranjo 2:1 (fatia permanente ≥ m·W/(1−m))',
              modelo=100 * m_max, referencia=f'Corredor: carregamento p5 = {100 * cl:.1f}% e p95 = {100 * ch:.1f}% do SIL', fonte='Cadastro ONS de linhas + fluxo verificado N–NE',
              leitura='O 2:1 é viável para toda a faixa do slider (20–60%)' if m_max >= 0.6 else 'Limita o slider de mínimo de regulação'))
V.append(dict(id='eta', cenario='todos', verificacao='Histerese das manobras (eta, fração do degrau) – premissa, não calibrada', modelo=ETA0,
              referencia='Sensibilidade 0,1–0,75 em ModeloFisico_Taxas', fonte='Premissa declarada', leitura='O Centro Urbano reproduz 1 ciclo/dia para eta ≤ 0,5 (a forma do dia domina)'))
VAL = pd.DataFrame(V); VAL['modelo'] = VAL.modelo.astype(float).round(3)
VAL.to_csv(P + 'modelo_fisico_validacao.csv', index=False, sep=';', encoding='utf-8-sig')
print(VAL[['id', 'modelo', 'leitura']].to_string())

# ======================= 5. JSON (camada 1) =======================
J = dict(meta=dict(gerado_em='2026-10-13', script='scripts/09_modelo_fisico.py', descricao=__doc__.strip().split('\n\n')[0],
                   caso_base=dict(arranjo=ARR0, N=N_BASE, m=M0, k=K0, eta=ETA0, posicoes_tap=T_POS, banda_tap_passos=DB_TAP)),
         corredor_NNE=dict(circuitos=len(COR) - 1, km=float(COR.km.iloc[-1]), Qc_Mvar=float(QC), SIL_MW=float(SIL), carga_p5_SIL=cl, carga_p95_SIL=ch, m_max_2para1=m_max,
                           W_Mvar=float(W_campo), N_degraus=N_campo),
         series=prep, taxas=TX.to_dict(orient='records'), mensal=MM.to_dict(orient='records'), dia_tipico=tip, validacao=VAL.to_dict(orient='records'))
json.dump(J, open(P + 'modelo_fisico_camada1.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'), default=float)

# ======================= 6. Planilha =======================
F = 'Arial'
fT = Font(name=F, size=14, bold=True); fS = Font(name=F, size=9, italic=True, color='555555')
fH = Font(name=F, size=10, bold=True, color='FFFFFF'); fN = Font(name=F, size=10)
fBlue = Font(name=F, size=10, color='0000FF'); fGreen = Font(name=F, size=10, color='008000'); fBold = Font(name=F, size=10, bold=True)
hdrFill = PatternFill('solid', fgColor='1F3864'); keyFill = PatternFill('solid', fgColor='FFFF00'); thin = Side(style='thin', color='BFBFBF')
wb = openpyxl.load_workbook(XL)
NEW = ['Corredor_NNE', 'ModeloFisico_Taxas', 'ModeloFisico_Mensal', 'ModeloFisico_Validacao']
for nm in NEW:
    if nm in wb.sheetnames: del wb[nm]
pos = wb.sheetnames.index('Fontes_Reforcadas') if 'Fontes_Reforcadas' in wb.sheetnames else wb.sheetnames.index('Fontes_Primarias')
REF = {}
def sheet_df(name, idx, title, source, df, key=None, numfmt=None, widths=None, notes=None):
    ws = wb.create_sheet(name, idx)
    ws['A1'] = title; ws['A1'].font = fT
    ws['A2'] = 'Fonte: ' + source; ws['A2'].font = fS
    if notes: ws['A3'] = notes; ws['A3'].font = fS
    r0 = 4
    for jj, c in enumerate(df.columns, 1):
        cell = ws.cell(r0, jj, str(c)); cell.font = fH; cell.fill = hdrFill; cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for jj, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and np.isnan(v)): v = None
            elif isinstance(v, np.integer): v = int(v)
            elif isinstance(v, np.floating): v = float(v)
            cell = ws.cell(i, jj, v); cell.font = fN; cell.alignment = Alignment(wrap_text=True, vertical='top')
            if numfmt and df.columns[jj - 1] in numfmt: cell.number_format = numfmt[df.columns[jj - 1]]
        if key is not None:
            k = row[list(df.columns).index(key)]
            for jj, c in enumerate(df.columns, 1): REF[(name, k, c)] = f"'{name}'!${get_column_letter(jj)}${i}"
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    for jj, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(jj)].width = (widths or {}).get(c) or min(max(len(str(c)) + 2, 11), 40)
    ws.row_dimensions[r0].height = 30
    return ws
SRC9 = 'Script: scripts/09_modelo_fisico.py (camada 1 do motor; decisão P1). Dados: ONS cadastro de linhas, BALANCO_DESSEM_DETALHE e intercâmbio entre subsistemas.'
NOTE9 = ('r(t) = (u_hi² − u²)/(u_hi² − u_lo²), u = |fluxo ou carga líquida|, u_lo/u_hi = p5/p95: necessidade da compensação manobrável (sobra de reativo da linha Qc·(1 − (P/SIL)²)). '
         'Fixo: N degraus, manobra no meio do degrau ± histerese (eta × degrau). VSR 1:1 = tamanho W, mínimo m, sai quando r < m/2. '
         'VSR 2:1 = tamanho W/(1−m), substitui o banco manobrável e uma fatia da compensação permanente (CIGRÉ A2-211), sai só quando r < −m/(2(1−m)). '
         'k = escala da volatilidade (P4) em torno do perfil típico. Taps/dia são piso (resolução de 30 min a 1 h). Desajuste = |necessidade − compensação|, em unidades de W.')
sheet_df('Corredor_NNE', pos, 'Corredor Norte–Nordeste: linhas de fronteira, carregamento capacitivo e capacidade natural (SIL)',
         'ONS – cadastro de linhas de transmissão (LINHA_TRANSMISSAO.csv), linhas ativas com um terminal no N e outro no NE. SIL = 100·√(B/X), X e B em % na base 100 MVA. ' + SRC9.split('. Dados')[0],
         COR, key='linha', notes='O corredor também passa por subestações-polo (Miracema, Colinas, P. Dutra, Teresina II) que servem outras linhas; por isso a compensação permanente do corredor não é isolável no cadastro e o modelo usa a regra de dimensionamento por p5/p95.',
         numfmt={'km': '#,##0', 'X_pct_100MVA': '0.000', 'Qc_Mvar': '#,##0.0', 'SIL_MW': '#,##0', 'Mvar_km': '0.000', 'capacidade_longa_A': '#,##0'},
         widths={'linha': 62, 'entrada_operacao': 14})
sheet_df('ModeloFisico_Taxas', pos + 1, 'Modelo físico: manobras e taps por dia, por cenário e parâmetro (grade de sensibilidade)', SRC9, TX, key='id', notes=NOTE9,
         numfmt={c: '0.00' for c in TX.columns[8:]} | {'desajuste_fixo': '0.000', 'desajuste_vsr': '0.000', 'vsr_taps_dia': '0.0'}, widths={'id': 34})
sheet_df('ModeloFisico_Mensal', pos + 2, f'Modelo físico: caso-base por mês (arranjo 2:1, m = 50%, k = 1, eta = 0,3; N = {N_campo} no Campo e 1 nos demais)', SRC9, MM, notes=NOTE9,
         numfmt={c: '0.00' for c in MM.columns[4:]} | {'vsr_taps_dia': '0.0'})
sheet_df('ModeloFisico_Validacao', pos + 3, 'Modelo físico: validação contra a operação observada', SRC9, VAL, key='id',
         notes='Nenhum parâmetro foi ajustado para reproduzir as referências; eta é premissa declarada e a sensibilidade está na grade.',
         numfmt={'modelo': '0.00'}, widths={'verificacao': 50, 'referencia': 46, 'fonte': 50, 'leitura': 44})
# ---- Parametros_Modelo: P79+ ----
ws = wb['Parametros_Modelo']
first = next((r for r in range(5, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value[1:].isdigit() and int(ws.cell(r, 1).value[1:]) >= 79), None)
if first: ws.delete_rows(first, ws.max_row - first + 1)
TXR = lambda cen, arr, col, N_=None: '=' + REF[('ModeloFisico_Taxas', f'{cen}|{arr}|N{N_BASE[cen] if N_ is None else N_}|m{M0}|k{K0}|eta{ETA0}', col)]
PR = []
for cen in SER:
    nm = NOME[cen]
    PR += [('Modelo físico', f'Manobras/dia por disjuntor manobrável do fixo – {nm} (caso-base)', TXR(cen, '2:1', 'fixo_ops_dia_por_disjuntor_manobravel'), 'por dia', 'Desgaste do disjuntor do fixo = manobras evitadas por VSR', 'Script 09', 'ModeloFisico_Taxas', f'Arranjo 2:1, m = 50%, N = {N_BASE[cen]}.', False),
           ('Modelo físico', f'Manobras residuais/dia do VSR – {nm} (caso-base, 2:1)', TXR(cen, '2:1', 'vsr_manobras_residuais_dia'), 'por dia', 'Desgaste do disjuntor do VSR', 'Script 09', 'ModeloFisico_Taxas', 'Só em carregamento extremo.', False),
           ('Modelo físico', f'Operações de tap/dia do VSR – {nm} (caso-base)', TXR(cen, '2:1', 'vsr_taps_dia'), 'por dia', 'OPEX do OLTC', 'Script 09', 'ModeloFisico_Taxas', 'Piso: resolução de 30 min (DESSEM) ou 1 h (fluxo).', False),
           ('Modelo físico', f'Manobras residuais/dia do VSR no arranjo 1:1 – {nm}', TXR(cen, '1:1', 'vsr_manobras_residuais_dia', 1), 'por dia', 'Risco do caso 1:1 (2.4)', 'Script 09', 'ModeloFisico_Taxas', 'VSR do tamanho do banco manobrável (N = 1): sai quando a necessidade cai abaixo de m/2.', False)]
PR += [('Corredor N–NE', 'SIL total do corredor N–NE', '=' + REF[('Corredor_NNE', 'TOTAL', 'SIL_MW')], 'MW', 'Converte fluxo em necessidade de reator (Campo)', 'ONS – cadastro de linhas', 'Corredor_NNE', '11 circuitos (8 × 500 kV, 3 × 230 kV).', False),
       ('Corredor N–NE', 'Carregamento capacitivo total do corredor N–NE', '=' + REF[('Corredor_NNE', 'TOTAL', 'Qc_Mvar')], 'Mvar', 'Escala do slider de extensão (Campo)', 'ONS – cadastro de linhas', 'Corredor_NNE', '', False),
       ('Corredor N–NE', 'Degraus manobráveis do corredor N–NE (variação da sobra de reativo ÷ 150 Mvar)', '=' + REF[('ModeloFisico_Validacao', 'campo_N', 'modelo')], 'degraus', 'N do caso-base do Campo', 'Script 09', 'ModeloFisico_Validacao', '', False),
       ('Corredor N–NE', 'Mínimo de regulação máximo viável no arranjo 2:1 – Campo', '=' + REF[('ModeloFisico_Validacao', 'campo_m_max', 'modelo')], '% (escala 0–100)', 'Limite do slider de mínimo de regulação', 'Script 09', 'ModeloFisico_Validacao', 'Acima disso, a compensação permanente do corredor não comporta a fatia que o VSR substitui.', False),
       ('Modelo físico', 'Histerese das manobras (eta, fração do degrau)', ETA0, 'fração do degrau', 'Faixa que o operador tolera antes de manobrar', 'Premissa declarada', 'ModeloFisico_Taxas', 'Sensibilidade 0,1–0,75 na grade.', True)]
start = ws.max_row + 1
NF = {'adim.': '0.000', 'por dia': '0.00', 'MW': '#,##0', 'Mvar': '#,##0', 'fração do degrau': '0.00', 'degraus': '0'}
for i, (gr, pa, v, u, uso, fonte, aba, obs, key_) in enumerate(PR, start):
    vals = [f'P{79 + i - start}', gr, pa, v, u, uso, fonte, aba, obs]
    for jj, x in enumerate(vals, 1):
        c = ws.cell(i, jj, x); c.font = fN; c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(bottom=thin)
    vc = ws.cell(i, 4)
    vc.font = (fGreen if (v.count('!') == 1 and '(' not in v and '/' not in v) else fN) if isinstance(v, str) and v.startswith('=') else fBlue
    vc.number_format = NF.get(u, '#,##0.00')
    if key_: vc.fill = keyFill
# ---- LEIA-ME ----
ws = wb['LEIA-ME']
lines = []
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v == 'Abas da planilha': break
    lines.append([v or '', 'T' if r == 1 else ('B' if ws.cell(r, 1).font.b else 'N')])
lines = [x for x in lines if not x[0].startswith(('5. ModeloFisico', 'Necessidade normalizada r(t)'))]
out = []
for t, f in lines:
    if t.startswith('3. Reprodutibilidade'):
        t = '3. Reprodutibilidade: os arquivos brutos estão em raw/ e os scripts Python em scripts/ (00 a 09). Rodar na ordem regenera tudo; o 08 e o 09 rodam depois do 05, nessa ordem, e acrescentam as abas da v3. O manifesto traz o SHA-256 de cada arquivo bruto.'
    out.append((t, f))
    if t.startswith('4. Pesquisa_Literatura'):
        out.append(('5. ModeloFisico_* e Corredor_NNE: camada 1 do motor (script 09) – manobras do reator fixo, manobras residuais e taps do VSR por dia, por cenário, com grade de sensibilidade e validação contra a operação observada. As linhas P79 em diante de Parametros_Modelo puxam essas abas.', 'N'))
    if t.startswith('Volatilidade (decisão P4)'):
        out.append(('Necessidade normalizada r(t) = (u_hi² − u²)/(u_hi² − u_lo²), com u = |fluxo ou carga líquida| e u_lo/u_hi = p5/p95: fração da compensação manobrável necessária, derivada da sobra de reativo da linha Qc·(1 − (P/SIL)²).', 'N'))
ws.delete_rows(1, ws.max_row)
FM_ = {'T': fT, 'B': fBold, 'N': fN}
for i, (t, f) in enumerate(out, 1):
    c = ws.cell(i, 1, t if t else None); c.font = FM_[f]; c.alignment = Alignment(wrap_text=True, vertical='top')
r = len(out) + 2
ws.cell(r, 1, 'Abas da planilha').font = fBold
for k, name in enumerate(wb.sheetnames, r + 1): ws.cell(k, 1, name).font = fN
wb.save(XL)
pm = wb['Parametros_Modelo']
print('ok', XL, len(wb.sheetnames), 'abas; Parametros_Modelo até', pm.cell(pm.max_row, 1).value, f'({time.time() - t0:.0f}s)')
