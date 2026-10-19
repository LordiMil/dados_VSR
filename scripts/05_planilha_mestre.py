"""Monta a planilha-mestre PoliVSR_Base_de_Dados.xlsx a partir de processed/ e FONTES_manifesto.csv.
Uso: python 05_planilha_mestre.py <pasta_base_PoliVSR-dados> <arquivo_saida.xlsx>"""
import sys, pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
BASE, OUT = sys.argv[1], sys.argv[2]; P = BASE + '/processed/'
rd = lambda f: pd.read_csv(P + f, sep=';', encoding='utf-8-sig')
F = 'Arial'
fT = Font(name=F, size=14, bold=True); fS = Font(name=F, size=9, italic=True, color='555555')
fH = Font(name=F, size=10, bold=True, color='FFFFFF'); fN = Font(name=F, size=10)
fBlue = Font(name=F, size=10, color='0000FF'); fGreen = Font(name=F, size=10, color='008000'); fBold = Font(name=F, size=10, bold=True)
hdrFill = PatternFill('solid', fgColor='1F3864'); keyFill = PatternFill('solid', fgColor='FFFF00'); grpFill = PatternFill('solid', fgColor='D9E1F2')
thin = Side(style='thin', color='BFBFBF')
wb = Workbook(); wb.remove(wb.active)
REF = {}  # (sheet, key, col) -> address

def sheet_df(name, title, source, df, key=None, numfmt=None, widths=None, notes=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = fT
    ws['A2'] = 'Fonte: ' + source; ws['A2'].font = fS
    r0 = 4
    if notes:
        ws['A3'] = notes; ws['A3'].font = fS
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(r0, j, str(c)); cell.font = fH; cell.fill = hdrFill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (bool, np.bool_)): v = 'sim' if v else 'não'
            elif isinstance(v, float) and np.isnan(v): v = None
            elif isinstance(v, (np.integer,)): v = int(v)
            elif isinstance(v, (np.floating,)): v = float(v)
            elif isinstance(v, (bool, np.bool_)): v = 'sim' if v else 'não'
            cell = ws.cell(i, j, v); cell.font = fN
            if numfmt and df.columns[j - 1] in numfmt: cell.number_format = numfmt[df.columns[j - 1]]
        if key is not None:
            k = row[list(df.columns).index(key)] if isinstance(key, str) else tuple(row[list(df.columns).index(x)] for x in key)
            for j, c in enumerate(df.columns, 1): REF[(name, k, c)] = f"'{name}'!${get_column_letter(j)}${i}"
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c) or min(max(len(str(c)) + 2, 11), 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 30
    return ws, r0

ONS = 'ONS – Portal de Dados Abertos (dados.ons.org.br), acesso em 09/10/2026, licença cc-by. '
ANEEL = 'ANEEL – Portal de Dados Abertos (dadosabertos.aneel.gov.br), acesso em 09/10/2026, licença ODbL. '
BCB = 'Banco Central do Brasil – SGS (api.bcb.gov.br), acesso em 09/10/2026. '

# ---------------- data sheets ----------------
res = rd('resumo_01_reatores_linhas_disp.csv')
sheet_df('Resumo_Inventario', 'Resumo do inventário de reatores, linhas e disponibilidade', ONS + 'Conjuntos equipamento-controle-reativo, linha-transmissao, ind_disponibilidade_ft_reativo. Script: scripts/01_inventario_reatores_linhas.py',
         res, key='indicador', widths={'indicador': 62, 'nota': 70, 'secao': 16})
sub = rd('reatores_por_subsistema_tensao.csv')
sub.loc[len(sub)] = ['SIN (total)'] + list(sub.iloc[:, 1:].sum())
sheet_df('Reatores_Subsistema', 'Reatores shunt ativos ≥ 230 kV por subsistema e faixa de tensão', ONS + 'EQP_CONTROLE_REATIVO.csv (cadastro de equipamentos de controle de reativos da Rede de Operação).',
         sub, key='id_subsistema', numfmt={c: '#,##0' for c in sub.columns[1:]})
sheet_df('Reatores_UF', 'Reatores shunt ativos ≥ 230 kV por UF', ONS + 'EQP_CONTROLE_REATIVO.csv.', rd('reatores_por_UF.csv'), key='id_estado', numfmt={'Mvar': '#,##0'})
sheet_df('Reatores_Idade', 'Distribuição de idade dos reatores ativos ≥ 230 kV (em 09/10/2026)', ONS + 'EQP_CONTROLE_REATIVO.csv.',
         rd('reatores_distribuicao_idade.csv'), key='faixa_idade_anos', widths={'faixa_idade_anos': 36},
         notes='Atenção: 185 reatores têm data de entrada 1990-12-31, uma data-padrão do cadastro para equipamentos anteriores a 1991; a idade real deles é ≥ 35,8 anos.')
sa = rd('reatores_serie_anual_entradas.csv')
sheet_df('Reatores_SerieAnual', 'Série anual de entradas e desativações de reatores ≥ 230 kV', ONS + 'EQP_CONTROLE_REATIVO.csv (datas de entrada e desativação).',
         sa, key='ano', numfmt={'Mvar_entrada': '#,##0', 'entradas': '0', 'desativacoes': '0', 'estoque_fim_ano': '#,##0'},
         notes='Estoque reconstruído só a partir do cadastro atual; equipamentos pré-1991 aparecem todos em 1990 (data-padrão).')
la = rd('reatores_ativos_ONS.csv')
sheet_df('Reatores_Lista', 'Lista de reatores ativos (todas as tensões) – 1 linha por equipamento', ONS + 'EQP_CONTROLE_REATIVO.csv; colunas kv, faixa_kv, idade_anos e data_padrao_pre1991 derivadas.',
         la, widths={'nom_equipamento': 48, 'nom_subestacao': 24, 'nom_agente_proprietario': 28})
li = rd('linhas_carregamento_capacitivo_por_tensao.csv')
sheet_df('Linhas_Mvar_km', 'Carregamento capacitivo das linhas CA ativas (Mvar/km) por nível de tensão', ONS + 'LINHA_TRANSMISSAO.csv; Mvar/km = val_shunt ÷ val_comprimento (val_shunt = susceptância de sequência positiva em Mvar, convenção ANAREDE).',
         li, key='val_niveltensao_kv', numfmt={'km_total': '#,##0', 'Mvar_carregamento_total': '#,##0'})
ds = rd('disponibilidade_FTRT_estatisticas.csv')
sheet_df('Disp_FTRT_Estat', 'Disponibilidade mensal das Funções Transmissão de Reator (FTRT), estatísticas 2011–2026', ONS + 'IND_DISPONIBILIDADE_FT_REATIVO_2011..2026.csv (indicador DISP, Submódulo 9.2 dos Procedimentos de Rede).',
         ds, key='caracteristica', notes='indisp_horas_mes = (100 − DISP)/100 × horas do mês. Agregado por nível de tensão, não por equipamento.')
sheet_df('Disp_FTRT_Anual', 'Disponibilidade média anual das FTRT (%)', ONS + 'IND_DISPONIBILIDADE_FT_REATIVO_*.csv.', rd('disponibilidade_FTRT_media_anual.csv'), key='ano')
dm = rd('disponibilidade_FT_reativo_mensal.csv')
sheet_df('Disp_FT_Mensal', 'Série mensal de disponibilidade de todas as FT de reativo (reator, capacitor, compensadores)', ONS + 'IND_DISPONIBILIDADE_FT_REATIVO_*.csv.', dm,
         numfmt={'indisp_pct': '0.0000', 'horas_indisp_mes': '0.00'})
st = rd('dessem_estatisticas_curva_pato.csv')
sheet_df('CurvaPato_Estat', 'Curva de pato: estatísticas diárias da carga líquida por subsistema (mai/2025–out/2026)', ONS + 'BALANCO_DESSEM_DETALHE (499 arquivos diários, semi-horários). Carga líquida = demanda − (eólica + fotovoltaica + MMGD). Script: scripts/02_series_carga_curva_pato.py',
         st, key='cod_subsistema', numfmt={c: '#,##0' for c in st.columns if c.startswith(('amplitude', 'rampa', 'vale', 'liquida'))},
         notes='Valores em MW. Dados programados (DESSEM), não verificados. SIN = soma dos 4 subsistemas nos 465 dias completos. Carga líquida do NE pode ser negativa (exportador).')
sheet_df('CurvaPato_Mensal', 'Curva de pato: médias mensais por subsistema', ONS + 'BALANCO_DESSEM_DETALHE.', rd('dessem_mensal_por_subsistema.csv'), key=('cod_subsistema', 'mes'))
cpd = rd('dessem_diario_curva_pato_por_subsistema.csv')
sheet_df('CurvaPato_Diario', 'Curva de pato: métricas diárias por subsistema (série para o modelo estocástico)', ONS + 'BALANCO_DESSEM_DETALHE.', cpd,
         numfmt={c: '#,##0' for c in cpd.columns if c not in ('cod_subsistema', 'data', 'hora_liquida_min', 'hora_liquida_max', 'fracao_vre_no_vale', 'dia_semana', 'fim_de_semana', 'mes')})
sheet_df('Perfil_Mes_Hora', 'Perfil médio da carga líquida por mês e meia-hora (MW) – as curvas de pato', ONS + 'BALANCO_DESSEM_DETALHE.', rd('dessem_perfil_medio_mes_hora.csv'))
ca = rd('carga_horaria_anual_2015_2026.csv')
sheet_df('Carga_Anual', 'Carga horária 2015–2026: estatísticas anuais por subsistema', ONS + 'CURVA_CARGA_2015..2026.csv (carga de energia horária, MWmed). Script: scripts/02_series_carga_curva_pato.py',
         ca, key=('id_subsistema', 'ano'), numfmt={c: '#,##0' for c in ['carga_media', 'amplitude_media', 'amplitude_p95', 'rampa3h_media']},
         notes='2026 parcial (até 07/10/2026). fator_carga = média/máximo diário. pct_dias_min_10a15h = % de dias cujo mínimo horário ocorre entre 10h e 15h.')
sheet_df('Carga_Diaria', 'Carga horária 2015–2026: métricas diárias por subsistema', ONS + 'CURVA_CARGA_2015..2026.csv.', rd('carga_horaria_diario_2015_2026.csv'))
ma = rd('mmgd_aneel_anual.csv').rename(columns={'mes': 'ano'}); ma = ma[ma.ano >= 2012]
sheet_df('MMGD_Anual', 'MMGD: potência adicionada e acumulada por ano (GW)', ANEEL + 'Relação de empreendimentos de geração distribuída (período de referência 10/2026). Script: scripts/03_mmgd_aneel.py',
         ma, key='ano', notes='Data usada: DthAtualizaCadastralEmpreend (única data da base; proxy da data de conexão). Meses finais podem estar incompletos por atraso de registro. Registros com data 1900 (0,001 GW) excluídos.')
mm = rd('mmgd_aneel_GW_acumulado_mensal_regiao.csv'); mm = mm[mm.mes >= '2012-01']
sheet_df('MMGD_Mensal_Regiao', 'MMGD: estoque acumulado mensal por região (GW)', ANEEL + 'Relação de empreendimentos de geração distribuída.', mm, key='mes')
sheet_df('MMGD_UF', 'MMGD: estoque por UF', ANEEL + 'Relação de empreendimentos de geração distribuída.', rd('mmgd_aneel_estoque_por_UF.csv'), key='SigUF', numfmt={'qtd': '#,##0', 'GW': '0.000'})
sheet_df('MMGD_Distribuidora', 'MMGD: estoque por distribuidora', ANEEL + 'Relação de empreendimentos de geração distribuída.', rd('mmgd_aneel_estoque_por_distribuidora.csv'), key='SigAgente', numfmt={'qtd': '#,##0', 'GW': '0.000'})
mac = rd('macro_bcb_mensal.csv'); mac = mac[mac.ipca_mes_pct.notna()]
wsM, rM = sheet_df('Macro_Mensal', 'Séries macroeconômicas mensais (IPCA, Selic, IGP-M, câmbio)', BCB + 'SGS 433 (IPCA), 4390 (Selic acumulada no mês), 189 (IGP-M), 3698 (US$ venda, média mensal), 21619 (€ venda, diário → média mensal). Script: scripts/04_custos_bpr_macro.py',
         mac, key='mes', notes='ipca_indice_base_dez2009 = produto acumulado de (1 + IPCA mensal), base dez/2009 = 1. Taxas em % (não em fração).')
lastM = mac.mes.iloc[-1]
# BPR com fórmula de correção
dj = rd('bpr_disjuntor_preco_unitario.csv').drop_duplicates('NomTipoTensao').sort_values('kv')
dj['IndiceReferencia'] = dj.IndiceReferencia.str.strip().apply(lambda s: f"{s.split('/')[1]}-{int(s.split('/')[0]):02d}")
dj = dj[['kv', 'IndiceReferencia', 'preco']].rename(columns={'kv': 'tensao_kV', 'IndiceReferencia': 'mes_base_preco', 'preco': 'preco_base_R$'})
wsD, rD = sheet_df('BPR_Disjuntor', 'Preço unitário de referência do disjuntor (ANEEL BPR) e correção pelo IPCA', ANEEL + 'BPR – Banco de Preços de Referência (subestação), item-modulo-manobra, item "Disjuntor", classe UAR. Correção: fórmula com o índice IPCA da aba Macro_Mensal.',
         dj, key='tensao_kV', numfmt={'preco_base_R$': '#,##0.00'})
n = len(dj); mcol = list(mac.columns); ci = get_column_letter(mcol.index('ipca_indice_base_dez2009') + 1); cm = get_column_letter(mcol.index('mes') + 1)
rng_m = f"Macro_Mensal!${cm}${rM+1}:${cm}${rM+len(mac)}"; rng_i = f"Macro_Mensal!${ci}${rM+1}:${ci}${rM+len(mac)}"
wsD.cell(rD, 4, f'fator_IPCA_ate_{lastM}').font = fH; wsD.cell(rD, 4).fill = hdrFill
wsD.cell(rD, 5, f'preco_corrigido_R$_{lastM}').font = fH; wsD.cell(rD, 5).fill = hdrFill
wsD.cell(rD - 1, 4, 'mês final da correção:').font = fBold; wsD.cell(rD - 1, 5, lastM).font = fBlue
for i in range(rD + 1, rD + 1 + n):
    wsD.cell(i, 4, f'=INDEX({rng_i},MATCH($E${rD-1},{rng_m},0))/INDEX({rng_i},MATCH(B{i},{rng_m},0))').number_format = '0.0000'
    wsD.cell(i, 5, f'=C{i}*D{i}').number_format = '#,##0.00'
    REF[('BPR_Disjuntor', dj.tensao_kV.iloc[i - rD - 1], 'corrigido')] = f"'BPR_Disjuntor'!$E${i}"
    for c in (4, 5): wsD.cell(i, c).font = fN
wsD.column_dimensions['D'].width = 18; wsD.column_dimensions['E'].width = 24
sheet_df('BPR_Modulos_Reator', 'Custo total dos módulos de manobra (ANEEL BPR): conexão de reator de barra (CRB), de linha manobrável (CRL) e não manobrável (CRLn), e entrada de linha (EL)', ANEEL + 'BPR – item-modulo-manobra; custo_total = Σ(preço × quantidade) dos itens do módulo.',
         rd('bpr_modulos_manobra_reator_custo.csv'), numfmt={'custo_total': '#,##0'},
         notes='Valores nominais na base de preço de cada item (mai/2017 para equipamentos; dez/2011 para obras civis). Não corrigidos. O preço do reator em si não consta deste conjunto aberto.')
sheet_df('BPR_CRB_Disjuntor', 'Participação do disjuntor no custo do módulo de conexão de reator de barra (CRB)', ANEEL + 'BPR – item-modulo-manobra.',
         rd('bpr_CRB_participacao_disjuntor.csv'), key=('NomTipoTensao', 'DscTipoArranjo'), numfmt={'custo_modulo': '#,##0', 'custo_disjuntores': '#,##0'})
sheet_df('SIGET_RAP_Reatores', 'Receita Anual Permitida (RAP) por Função Transmissão de reator – estatísticas por tipo e tensão', ANEEL + 'SIGET – Lista de Módulos Prévia – Reajuste RAP (situação "Ativa"; RAP do ciclo = soma dos módulos da FT). Script 07.',
         rd('siget_rap_reatores_estatisticas.csv'), key=('tipo', 'kV'), numfmt={c: '#,##0' for c in ['p10', 'p25', 'mediana', 'p75', 'p90', 'media', 'mediana_RAP_por_Mvar']},
         notes='RTB = reator de barra; RTL = reator de linha. PB (pagamento-base mensal) = RAP ÷ 12. Lista completa por FT em processed/siget_rap_FT_reatores.csv.')
sheet_df('PVI_Reator', 'Parcela Variável por Indisponibilidade (PVI) de uma FT de reator, por evento', 'ANEEL – Regras dos Serviços de Transmissão, Módulo 4 rev.2 (REN 1.137/2025), Seção 4.3 item 4.2 Eq.3 e Anexo I (REA: Ko = 150, Kp = 10). Script 07.',
         rd('pvi_reator_por_evento.csv'), key=('kV', 'tipo', 'duracao_min'), numfmt={c: '#,##0' for c in ['RAP_FT_RTB_mediana_RS', 'PB_mensal_RS', 'PVI_RS', 'parcela_no_mes_RS', 'diferido_meses_seguintes_RS', 'teto_12m_RS', 'PVI_RS_apos_teto_12m']},
         notes='Mês de 30 dias (43.200 min). Ko vale nos primeiros 300 min de um desligamento forçado e cai para Kp depois. Limite: 50% do PB no mês (saldo diferido) e 25% de 12 PB em 12 meses.')
sheet_df('Fontes_Reforcadas', 'Fatos documentais com página/seção (atos ANEEL, ONS PAR/PEL 2025, literatura VSR)', 'Script 07 – processed/fontes_reforcadas.csv. Páginas do PAR/PEL localizadas automaticamente no texto extraído do PDF.',
         rd('fontes_reforcadas.csv'), widths={'parametro': 34, 'valor': 60, 'uso_no_modelo': 36, 'fonte': 40, 'arquivo_local': 44, 'localizacao': 28, 'url': 60, 'nivel_fonte': 24})
# ---------------- Fontes ----------------
man = pd.read_csv(BASE + '/FONTES_manifesto.csv', sep=';', encoding='utf-8-sig')
sheet_df('Fontes_Primarias', 'Catálogo de fontes primárias (arquivos brutos baixados)', 'Gerado por scripts/00_manifesto_fontes.py; SHA-256 permite verificar que o arquivo não foi alterado.',
         man, widths={'arquivo_local': 60, 'url_origem': 70, 'pagina_portal': 50, 'sha256': 66, 'conjunto': 34, 'instituicao': 30})
sec = pd.DataFrame([
 ['S01', 'CIGRÉ TB 510 (2012) – pesquisa internacional de confiabilidade de disjuntores AT', 'https://www.bib.irb.hr/599184', 'Taxa de falha grave 0,30/100 disjuntores-ano; manobra de reator shunt tem a maior taxa', 'B – técnico revisado'],
 ['S02', 'CIGRÉ SC A3 – Reliability Survey on Equipment', 'https://cse.cigre.org/cse-n023/cigre-reliability-survey-on-equipment', 'Faixa entre países 0,06 a 2,92 /100 disjuntores-ano; manobra frequente de reator eleva falhas', 'B'],
 ['S03', 'IEEE PES Switchgear Committee – TF Shunt Reactor Switching Endurance, ata 11/10/2017', 'https://bibliotheek.ehb.be:3193/soc/pes/switchgear/minutes/2017-2/F17HVCBa13REV0.pdf', '"Shunt reactors typically fail around 500–1000 operations"; endurance real ≈ 5× menor que a prevista; classe M2 recomendada', 'B'],
 ['S04', 'IEC 62271-100 (classes de endurance mecânica), via IEEE PES', 'https://bibliotheek.ehb.be:3193/soc/pes/switchgear/minutes/2001-2/F01HVCBa6.pdf', 'M1 = 2.000 operações; M2 = 10.000 operações', 'A – norma'],
 ['S05', 'CIGRÉ A3-10747 (2026) – Shunt Reactor Switching in Brazilian Grid', 'https://www.e-cigre.org/publications/detail/a3-10747-2026-shunt-reactor-switching-in-brazilian-grid-a-critical-analysis-of-circuit-breaker-standard-requirements-and-the-actual-stresses-in-the-network.html', '≈ 1.100 bancos de reatores ≥ 230 kV, ≈ 530 manobráveis; estresse dielétrico maior que o dos ensaios IEC', 'B'],
 ['S06', 'Reinhausen – VACUTAP VR (folheto)', 'https://www.reinhausen.com/fileadmin/downloadcenter/products/oltc/vr/flyer/in5062036_en.pdf', 'OLTC a vácuo aplicável a VSR: 300.000 comutações sem manutenção', 'C – fabricante'],
 ['S07', 'Reinhausen – VACUTAP VM (via Foundry Magazine)', 'https://www.foundrymag.com/new-products/efficient-tap-changing-arc-melting', 'OLTC a óleo convencional: manutenção a cada 40.000–100.000 comutações', 'D – imprensa citando fabricante'],
 ['S08', 'Hitachi Energy – Variable Shunt Reactors (CIGRÉ A2-211, 2014)', 'https://www.hitachienergy.com/content/dam/pg/references/files/variable-shunt-reactors-cigre2014-a2-211-2014.pdf', 'Exemplo de faixa: 420 kV, 90–200 Mvar', 'C'],
 ['S09', 'ANEEL – WACC regulatório 2026 (via Cenário Energia, 05/03/2026)', 'https://cenarioenergia.com.br/2026/03/05/aneel-wacc-regulatorio-2026-remuneracao-investimento-setor-eletrico/', 'Transmissão: 8,00% real pós-impostos; 12,11% real pré-impostos; vigência a partir de 01/03/2026', 'D – superada: confirmada no Despacho ANEEL 675/2026 (aba Fontes_Reforcadas)'],
 ['S10', 'ANEEL – WACC regulatório 2026 (via Agência iNFRA, 05/03/2026)', 'https://agenciainfra.com/blog/aneel-revisa-taxa-de-remuneracao-do-capital-do-setor-eletrico/', 'Confirma 8,00% / 12,11% para transmissão', 'D – superada pelo Despacho 675/2026'],
 ['S11', 'Siemens – Variable Shunt Reactor, Amprion (2017)', 'https://assets.new.siemens.com/siemens/assets/api/uuid:5b992c0e14405f78b85c97c447d74f2d717f971e/version:1542790208/2017-01-27-amprion-vsr-final-en.pdf', '400 kV, 50–250 Mvar, faixa de 80%, 33 posições de tap; 1ª unidade 2016', 'C – fabricante'],
 ['S12', 'Bengtsson – Variable shunt reactors for network stability control (ABB Review, Special Report)', 'https://library.e.abb.com/public/b9950e8daa93c268c1257ab8003b1097/41-44%20sr109_72dpi.pdf', 'Disjuntores externos fazem menos operações e exigem menos manutenção; VSR 420 kV 120–200 → 90–200 Mvar', 'C – fabricante'],
 ['S13', 'Hitachi Energy – Ensuring power quality and reliability in Norway', 'https://www.hitachienergy.com/news-and-events/customer-stories/ensuring-power-quality-and-reliability-in-norway', '10 VSR 420 kV 90–200 Mvar para Statnett e BKK desde 2012', 'C – fabricante'],
 ['S14', 'Transgrid – Substations Renewal and Maintenance Strategy (nov/2021), publicado pelo AER', 'https://aer.gov.au/system/files/Transgrid%20-%20Substations%20Renewal%20and%20Maintenance%20Strategy%20-%2030%20Nov%202021%20-%20PUBLIC.pdf', 'Vida nominal de reator a óleo 30 anos (§5.4.1 p.58); reatores manobrados diariamente (p.57); disjuntores de reator com regime "very onerous" (p.26)', 'B – documento regulatório público'],
], columns=['id', 'referencia', 'url', 'dado_usado', 'nivel_confianca'])
sheet_df('Fontes_Secundarias', 'Fontes secundárias (literatura técnica, fabricantes, imprensa) usadas como parâmetros', 'Ver documento "PoliVSR — Base de Dados e Fontes do Modelo" para contexto e ressalvas.',
         sec, widths={'referencia': 58, 'url': 60, 'dado_usado': 70, 'nivel_confianca': 26})

# ---------------- Parametros_Modelo ----------------
ws = wb.create_sheet('Parametros_Modelo', 0)
ws['A1'] = 'PoliVSR – Parâmetros do modelo com fonte'; ws['A1'].font = fT
ws['A2'] = 'Legenda: verde = valor vindo de outra aba (dado processado) · preto = fórmula · azul = valor de literatura digitado (ver Fontes_Secundarias) · amarelo = premissa-chave a confirmar.'; ws['A2'].font = fS
hdr = ['id', 'grupo', 'parametro', 'valor', 'unidade', 'uso_no_modelo', 'fonte', 'aba_origem', 'observacao']
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.font = fH; c.fill = hdrFill
R = lambda s, k, c: '=' + REF[(s, k, c)]
rows = [
 ('Reatores', 'Reatores shunt ativos ≥ 230 kV', R('Resumo_Inventario', 'Reatores ativos ≥ 230 kV', 'valor'), 'un', 'Universo do dashboard (substitui o placeholder "13")', 'ONS – cadastro de equipamentos de controle de reativos', 'Resumo_Inventario', 'Bate com ≈ 1.100 do CIGRÉ A3-10747 (S05). O cadastro não informa quais são manobráveis.'),
 ('Reatores', 'Reatores manobráveis ≥ 230 kV (literatura)', 530, 'un', 'Universo de reatores sujeitos a manobra diária', 'CIGRÉ A3-10747 (2026)', 'Fontes_Secundarias', 'Único número público de manobráveis; usar como cenário central.'),
 ('Reatores', 'Potência total dos reatores ativos ≥ 230 kV', R('Resumo_Inventario', 'Potência total ativa ≥ 230 kV', 'valor'), 'Mvar', 'Escala do sistema', 'ONS – cadastro', 'Resumo_Inventario', ''),
 ('Reatores', 'Potência mediana por reator ≥ 230 kV', R('Resumo_Inventario', 'Potência mediana por reator (≥ 230 kV)', 'valor'), 'Mvar', 'Tamanho típico do reator nos cenários', 'ONS – cadastro', 'Resumo_Inventario', ''),
 ('Reatores', 'Reatores ≥ 230 kV anteriores a 1991', R('Reatores_Idade', 'pré-1991 (data-padrão 1990-12-31)', 'qtd_reatores'), 'un', 'Candidatos a substituição por fim de vida útil', 'ONS – cadastro', 'Reatores_Idade', 'Idade real ≥ 35,8 anos (data-padrão do cadastro).'),
 ('Reatores', 'Crescimento do estoque de reatores ≥ 230 kV, 2015→2026', f"={REF[('Reatores_SerieAnual', 2026, 'estoque_fim_ano')]}/{REF[('Reatores_SerieAnual', 2015, 'estoque_fim_ano')]}-1", '%', 'Tendência de demanda por compensação reativa', 'ONS – cadastro', 'Reatores_SerieAnual', 'Estoque reconstruído a partir do cadastro atual.'),
 ('Linhas', 'Carregamento capacitivo mediano – 230 kV', R('Linhas_Mvar_km', 230.0, 'Mvar_km_mediana'), 'Mvar/km', 'Física do slider "Extensão do trecho"', 'ONS – cadastro de linhas', 'Linhas_Mvar_km', 'p10–p90 na aba de origem.'),
 ('Linhas', 'Carregamento capacitivo mediano – 500 kV', R('Linhas_Mvar_km', 500.0, 'Mvar_km_mediana'), 'Mvar/km', 'Física do slider "Extensão do trecho"', 'ONS – cadastro de linhas', 'Linhas_Mvar_km', 'p10 1,17; p90 2,13 (linhas compactas/LPNE puxam a cauda).'),
 ('Linhas', 'Grau de compensação shunt sistêmico 440–525 kV', R('Resumo_Inventario', 'Grau de compensação shunt sistêmico 440-525 kV (reatores ÷ carregamento)', 'valor'), 'pu', 'Mvar de reator por Mvar de carregamento da linha', 'ONS – cadastros de reatores e linhas', 'Resumo_Inventario', 'Razão agregada do SIN, não projeto de uma linha específica.'),
 ('Disponibilidade', 'Disponibilidade média mensal FTRT 500 kV (2011–2026)', R('Disp_FTRT_Estat', 'FTRT - 500', 'disp_media'), '% (escala 0–100)', 'Base para indisponibilidade/PVI do reator fixo', 'ONS – indicador DISP', 'Disp_FTRT_Estat', ''),
 ('Disponibilidade', 'Indisponibilidade média FTRT 500 kV', R('Disp_FTRT_Estat', 'FTRT - 500', 'indisp_horas_mes_media'), 'h/mês', 'Horas fora de serviço por mês (média do nível de tensão)', 'ONS – indicador DISP', 'Disp_FTRT_Estat', 'Inclui desligamentos programados e forçados.'),
 ('Disponibilidade', 'Indisponibilidade p95 FTRT 500 kV', R('Disp_FTRT_Estat', 'FTRT - 500', 'indisp_horas_mes_p95'), 'h/mês', 'Cauda de risco', 'ONS – indicador DISP', 'Disp_FTRT_Estat', ''),
 ('Curva de pato', 'Amplitude diária média da carga líquida – SIN', R('CurvaPato_Estat', 'SIN', 'amplitude_media'), 'MW', 'Volatilidade-base do cenário nacional', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', 'mai/2025–out/2026; carga líquida = demanda − eólica − FV − MMGD.'),
 ('Curva de pato', 'Amplitude diária p95 – SIN', R('CurvaPato_Estat', 'SIN', 'amplitude_p95'), 'MW', 'Cenário pessimista', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Curva de pato', 'Rampa máxima de 3 h, média diária – SIN', R('CurvaPato_Estat', 'SIN', 'rampa3h_media'), 'MW', 'Velocidade de variação que o controle de tensão acompanha', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Curva de pato', 'Rampa máxima de 3 h no período – SIN', R('CurvaPato_Estat', 'SIN', 'rampa3h_max'), 'MW', 'Evento extremo', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Curva de pato', '% de dias com mínimo da carga líquida entre 10h e 15h – SIN', R('CurvaPato_Estat', 'SIN', 'pct_dias_min_entre_10e15h'), '% (escala 0–100)', 'Frequência do "vale do meio-dia"', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Curva de pato', 'Amplitude diária média – Nordeste', R('CurvaPato_Estat', 'NE', 'amplitude_media'), 'MW', 'Cenários Campo/Indústria', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', 'Desvio-padrão e percentis na aba de origem.'),
 ('Curva de pato', 'Coeficiente de variação da amplitude – Nordeste', R('CurvaPato_Estat', 'NE', 'cv_amplitude'), 'adim.', 'Dispersão para a simulação de Monte Carlo', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Curva de pato', 'Amplitude diária média – Sudeste/Centro-Oeste', R('CurvaPato_Estat', 'SE', 'amplitude_media'), 'MW', 'Cenário Centro Urbano', 'ONS – DESSEM detalhado', 'CurvaPato_Estat', ''),
 ('Carga histórica', 'Variação da amplitude diária média do SIN, 2019→2025', f"={REF[('Carga_Anual', ('SIN', 2025), 'amplitude_media')]}/{REF[('Carga_Anual', ('SIN', 2019), 'amplitude_media')]}-1", '%', 'Tendência da volatilidade', 'ONS – curva de carga horária', 'Carga_Anual', 'Anos completos; 2026 é parcial.'),
 ('Carga histórica', '% de dias com mínimo entre 10h e 15h – NE, 2015', R('Carga_Anual', ('NE', 2015), 'pct_dias_min_10a15h'), '% (escala 0–100)', 'Linha de base antes da MMGD', 'ONS – curva de carga horária', 'Carga_Anual', ''),
 ('Carga histórica', '% de dias com mínimo entre 10h e 15h – NE, 2026 (parcial)', R('Carga_Anual', ('NE', 2026), 'pct_dias_min_10a15h'), '% (escala 0–100)', 'Evidência da curva de pato no NE', 'ONS – curva de carga horária', 'Carga_Anual', 'Até 07/10/2026.'),
 ('MMGD', 'Estoque de MMGD – Brasil (último mês da base)', R('MMGD_Mensal_Regiao', mm.mes.iloc[-1], 'Brasil'), 'GW', 'Motor da volatilidade', 'ANEEL – empreendimentos de GD', 'MMGD_Mensal_Regiao', 'Últimos meses podem ter atraso de registro.'),
 ('MMGD', 'Estoque de MMGD – Nordeste', R('MMGD_Mensal_Regiao', mm.mes.iloc[-1], 'Nordeste'), 'GW', 'Calibrar volatilidade regional', 'ANEEL – empreendimentos de GD', 'MMGD_Mensal_Regiao', ''),
 ('MMGD', 'MMGD adicionada em 2025', R('MMGD_Anual', 2025, 'GW_adicionados'), 'GW', 'Ritmo de crescimento', 'ANEEL – empreendimentos de GD', 'MMGD_Anual', ''),
 ('Custos', 'Disjuntor 500 kV – preço de referência (base mai/2017)', R('BPR_Disjuntor', 500.0, 'preco_base_R$'), 'R$', 'Custo de troca do disjuntor', 'ANEEL – BPR', 'BPR_Disjuntor', 'Item "Disjuntor", classe UAR.'),
 ('Custos', 'Disjuntor 500 kV – corrigido pelo IPCA', '=' + REF[('BPR_Disjuntor', 500.0, 'corrigido')], 'R$', 'Custo de troca do disjuntor (valores atuais)', 'ANEEL – BPR + BCB IPCA', 'BPR_Disjuntor', 'Correção por fórmula; muda se o mês final em BPR_Disjuntor!E3 for alterado.'),
 ('Custos', 'Disjuntor 230 kV – corrigido pelo IPCA', '=' + REF[('BPR_Disjuntor', 230.0, 'corrigido')], 'R$', 'Custo de troca do disjuntor (cenários 230 kV)', 'ANEEL – BPR + BCB IPCA', 'BPR_Disjuntor', ''),
 ('Custos', 'Participação do disjuntor no módulo CRB 500 kV (arranjo DJM)', R('BPR_CRB_Disjuntor', (500, 'DJM'), 'pct_disjuntor'), '% (escala 0–100)', 'Peso do disjuntor no custo de conexão do reator', 'ANEEL – BPR', 'BPR_CRB_Disjuntor', 'Arranjos AN e BDDD na aba de origem (32% a 47%).'),
 ('Financeiro', 'WACC regulatório transmissão – real pós-impostos (2026)', 0.08, '%', 'Taxa de desconto do VPL/payback', 'ANEEL – Despacho nº 675/2026, Anexo 1', 'Fontes_Reforcadas', 'Ato primário (vigência 01/03/2026); Despacho 1.174/2026 não altera a transmissão.'),
 ('Financeiro', 'WACC regulatório transmissão – real pré-impostos (2026)', 0.1211, '%', 'Sensibilidade', 'ANEEL – Despacho nº 675/2026, Anexo 1', 'Fontes_Reforcadas', ''),
 ('Financeiro', 'IPCA acumulado em 12 meses (último mês)', R('Macro_Mensal', lastM, 'ipca_12m_pct'), '% (escala 0–100)', 'Correção de custos', 'BCB – SGS 433', 'Macro_Mensal', ''),
 ('Financeiro', 'Juro real ex-post 12 meses (Selic ÷ IPCA)', R('Macro_Mensal', lastM, 'juro_real_12m_pct'), '% (escala 0–100)', 'Custo de oportunidade alternativo', 'BCB – SGS 4390 e 433', 'Macro_Mensal', ''),
 ('Financeiro', 'Câmbio US$/R$ (média do último mês)', R('Macro_Mensal', lastM, 'usd_brl_media'), 'R$/US$', 'Converter preços de equipamento em dólar', 'BCB – SGS 3698', 'Macro_Mensal', ''),
 ('Financeiro', 'Câmbio €/R$ (média do último mês)', R('Macro_Mensal', lastM, 'eur_brl_media'), 'R$/€', 'Converter preços de equipamento em euro', 'BCB – SGS 21619', 'Macro_Mensal', ''),
 ('Confiabilidade (literatura)', 'Taxa de falha grave de disjuntores SF6 (média mundial)', 0.30, 'falhas/100 disjuntores-ano', 'Taxa-base do modelo de falhas', 'CIGRÉ TB 510 (S01)', 'Fontes_Secundarias', 'Manobra de reator tem a maior taxa entre as aplicações.'),
 ('Confiabilidade (literatura)', 'Taxa de falha – mínimo entre países', 0.06, 'falhas/100 disjuntores-ano', 'Limite inferior da distribuição', 'CIGRÉ SC A3 (S02)', 'Fontes_Secundarias', ''),
 ('Confiabilidade (literatura)', 'Taxa de falha – máximo entre países', 2.92, 'falhas/100 disjuntores-ano', 'Limite superior da distribuição', 'CIGRÉ SC A3 (S02)', 'Fontes_Secundarias', ''),
 ('Confiabilidade (literatura)', 'Vida em manobra de reator – limite inferior', 500, 'operações', 'Distribuição de vida (ex.: Weibull)', 'IEEE PES TF 2017 (S03)', 'Fontes_Secundarias', 'Confirmar se o dado é do disjuntor ou do conjunto disjuntor + reator.'),
 ('Confiabilidade (literatura)', 'Vida em manobra de reator – limite superior', 1000, 'operações', 'Distribuição de vida', 'IEEE PES TF 2017 (S03)', 'Fontes_Secundarias', ''),
 ('Confiabilidade (literatura)', 'Razão endurance prevista ÷ real', 5, 'x', 'Incerteza da vida útil (fan chart)', 'IEEE PES TF 2017 (S03)', 'Fontes_Secundarias', '"often not met in practice by a factor of 5 or so".'),
 ('Confiabilidade (literatura)', 'Endurance mecânica classe M2', 10000, 'operações', 'Teto mecânico do disjuntor', 'IEC 62271-100 (S04)', 'Fontes_Secundarias', ''),
 ('VSR (fabricante)', 'OLTC a vácuo – comutações sem manutenção', 300000, 'comutações', 'OPEX do VSR', 'Reinhausen (S06)', 'Fontes_Secundarias', 'Dado de fabricante.'),
 ('VSR (fabricante)', 'OLTC a óleo – intervalo de manutenção (limite inferior)', 40000, 'comutações', 'Cenário conservador para o VSR', 'Reinhausen via imprensa (S07)', 'Fontes_Secundarias', 'Limite superior 100.000.'),
 ('VSR (fabricante)', 'Posições de tap do VSR', 33, 'posições', 'Resolução do comutador (taps)', 'Siemens – VSR Amprion 400 kV (S11)', 'Fontes_Secundarias', 'Mais posições = mais comutações = mais desgaste do OLTC (conservador para o VSR).'),
 ('Parcela Variável', 'Fator Ko do reator (outros desligamentos)', 150, 'adim.', 'PVI de falha forçada', 'ANEEL – Regras de Transmissão, Módulo 4 rev.2, Anexo I', 'Fontes_Reforcadas', 'Vale nos primeiros 300 minutos.'),
 ('Parcela Variável', 'Fator Kp do reator (programado e após 300 min)', 10, 'adim.', 'PVI de desligamento programado', 'ANEEL – Regras de Transmissão, Módulo 4 rev.2, Anexo I', 'Fontes_Reforcadas', ''),
 ('Parcela Variável', 'PVI mínima de uma falha forçada ≥ 5 h, em múltiplos do PB', '=150*300/(24*60*30)', 'PB', 'Custo regulatório somado a cada falha grave', 'Módulo 4 rev.2, Eq. 3', 'PVI_Reator', 'Piso: só os 300 minutos com Ko; mês de 30 dias.'),
 ('Parcela Variável', 'RAP mediana de FT de reator de barra – 230 kV', R('SIGET_RAP_Reatores', ('RTB', 230.0), 'mediana'), 'R$/ano', 'PB = RAP ÷ 12 (cenário Urbano)', 'ANEEL – SIGET', 'SIGET_RAP_Reatores', 'Mediana de 20 Mvar por FT.'),
 ('Parcela Variável', 'RAP mediana de FT de reator de barra – 500 kV', R('SIGET_RAP_Reatores', ('RTB', 500.0), 'mediana'), 'R$/ano', 'PB = RAP ÷ 12 (cenários Campo e Indústria)', 'ANEEL – SIGET', 'SIGET_RAP_Reatores', 'Mediana de 136 Mvar por FT.'),
 ('Parcela Variável', 'PVI de uma falha forçada de 5 h – 500 kV', R('PVI_Reator', (500, 'forcado (OUTROS DESLIGAMENTOS)', 300), 'PVI_RS'), 'R$', 'Custo regulatório por falha (Campo/Indústria)', 'Módulo 4 rev.2 + SIGET', 'PVI_Reator', '24 h: ver aba de origem.'),
 ('Operação (ONS)', 'Reatores de barra chaveados por dia útil de verão – 2026', 140, 'un', 'Corrobora ~1 ciclo/dia por reator (~2 operações de disjuntor)', 'ONS – PAR/PEL 2025, p. 64', 'Fontes_Reforcadas', '16.400 Mvar; + ~60 capacitores.'),
 ('Operação (ONS)', 'Reatores de barra chaveados por dia útil de verão – 2030', 170, 'un', 'Tendência das manobras', 'ONS – PAR/PEL 2025, p. 64', 'Fontes_Reforcadas', '19.200 Mvar; manobras 200 → 240 (+20%).'),
 ('Vida útil', 'Vida nominal de reator a óleo', 30, 'anos', 'Horizonte da simulação', 'Transgrid 2021 via AER (S14)', 'Fontes_Secundarias', '§5.4.1, p. 58.'),
 ('Manutenção', 'Periodicidade máxima da preventiva de disjuntores', 72, 'meses', 'Piso regulatório de manutenção', 'ANEEL – Regras de Transmissão, Módulo 4 rev.2, Seção 4.2', 'Fontes_Reforcadas', 'Item 6.6: troca de contatos avaliada pelo nº de operações.'),
]
for i, (g, p, v, u, uso, fonte, aba, obs) in enumerate(rows, 5):
    vals = [f'P{i-4:02d}', g, p, v, u, uso, fonte, aba, obs]
    for j, x in enumerate(vals, 1):
        c = ws.cell(i, j, x); c.font = fN; c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=thin)
    vc = ws.cell(i, 4)
    if isinstance(v, str) and v.startswith('=') and v.count('!') == 1 and '/' not in v: vc.font = fGreen
    elif isinstance(v, str): vc.font = fN
    else: vc.font = fBlue
    if u == '%' and not (isinstance(v, str) and 'Disp' in v) and g in ('Reatores', 'Carga histórica', 'Financeiro') and ('-1' in str(v) or isinstance(v, float)):
        vc.number_format = '0.0%'
    elif u in ('R$', 'R$/ano'): vc.number_format = '#,##0'
    elif u in ('MW', 'Mvar', 'un', 'operações', 'comutações', 'posições', 'anos', 'meses'): vc.number_format = '#,##0'
    elif u == 'PB': vc.number_format = '0.0000'
    else: vc.number_format = '#,##0.00'
    if p.startswith(('Reatores manobráveis', 'Vida em manobra')): vc.fill = keyFill
for j, w in enumerate([6, 18, 52, 16, 20, 40, 34, 22, 52], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'D5'

# ---------------- LEIA-ME ----------------
ws = wb.create_sheet('LEIA-ME', 0)
txt = [
 ('PoliVSR – Base de Dados do Modelo (v2, 09/10/2026 – fontes reforçadas)', fT),
 ('Planilha-mestre com as séries temporais e os parâmetros do modelo estocástico da calculadora "O Custo do Fixo". Todo número tem fonte: primária (ONS, ANEEL, BCB) na aba Fontes_Primarias, ou secundária (CIGRÉ, IEEE, IEC, fabricantes) na aba Fontes_Secundarias.', fN),
 ('', fN),
 ('Como usar', fBold),
 ('1. Parametros_Modelo: a tabela que alimenta a calculadora. Cada linha diz o valor, a unidade, para que serve no modelo, a fonte e a aba de origem.', fN),
 ('2. Abas de dados: séries e estatísticas processadas. As linhas 1–3 de cada aba dizem de onde vêm os dados e o script que os gerou.', fN),
 ('3. Reprodutibilidade: os arquivos brutos estão em raw/ e os scripts Python em scripts/ (00 a 07). Rodar na ordem regenera tudo. O manifesto traz o SHA-256 de cada arquivo bruto.', fN),
 ('', fN),
 ('Definições', fBold),
 ('Carga líquida = demanda − (geração eólica + fotovoltaica centralizada + MMGD), por meia-hora, a partir do balanço DESSEM detalhado do ONS (valores programados).', fN),
 ('Amplitude diária = máximo − mínimo da carga líquida no dia. Rampa de 3 h = maior aumento da carga líquida em 6 meias-horas consecutivas.', fN),
 ('Disponibilidade FTRT = indicador DISP mensal do ONS para as Funções Transmissão de reator, agregado por nível de tensão.', fN),
 ('', fN),
 ('Limitações conhecidas', fBold),
 ('• O cadastro do ONS não distingue reatores manobráveis de fixos; o número de manobráveis (≈ 530) vem do CIGRÉ A3-10747.', fN),
 ('• 185 reatores têm data de entrada 1990-12-31 (data-padrão para equipamentos anteriores a 1991).', fN),
 ('• Não há dado público de manobras por dia por reator. O modelo deve derivá-lo da curva de carga líquida com uma regra declarada.', fN),
 ('• O preço do reator em si e o do VSR não constam do BPR aberto da ANEEL; o prêmio de CAPEX do VSR depende de dado da Hitachi Energy.', fN),
 ('• MMGD: a base da ANEEL só traz a data de atualização cadastral (proxy da conexão); os últimos meses podem estar incompletos.', fN),
 ('• WACC 2026 confirmado no ato primário (Despacho ANEEL 675/2026). A Parcela Variável segue o Módulo 4 rev.2 das Regras de Transmissão (REN 1.137/2025); a REN 729/2016 está revogada.', fN),
 ('• Sem fonte aberta para o custo de manutenção do disjuntor como % do preço: segue como premissa (10%, ajustável).', fN),
 ('', fN),
 ('Convenções de cor (Parametros_Modelo)', fBold),
 ('Verde = valor puxado de outra aba · Preto = fórmula · Azul = valor de literatura digitado · Fundo amarelo = premissa-chave a confirmar.', fN),
 ('', fN),
 ('Licenças', fBold),
 ('ONS: Creative Commons Atribuição (cc-by). ANEEL: Open Database License (ODbL). Citar a instituição e o conjunto de dados ao usar.', fN),
]
for i, (t, f) in enumerate(txt, 1):
    c = ws.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 140
r = len(txt) + 2
ws.cell(r, 1, 'Abas da planilha').font = fBold
for k, name in enumerate(wb.sheetnames, r + 1):
    ws.cell(k, 1, name).font = fN
wb.save(OUT); print('ok', OUT, len(wb.sheetnames), 'abas')
